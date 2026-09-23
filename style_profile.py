"""
style_profile.py — learns THIS resume's own formatting conventions from
its own untouched content, so anything the pipeline writes or rewrites
(a replaced project, a recovered section, a re-tailored bullet) matches
whatever this specific resume already does — never a hardcoded
assumption like "projects have a tech-stack line in parens and a
GitHub link."

Two resumes can legitimately look completely different: one might
bracket its tech stack in (), another in [], another not at all;
one might bold technology keywords inline in bullets, another never
does; one might link every project, another links none. None of these
are "the right way" — they're just this document's own style, and the
only correct move is to detect it and reuse it, not to bake in an
opinion about what a resume "should" have.

This is READ-ONLY analysis over paragraphs docx_sections.py already
extracted — it doesn't touch the document, and it never needs the LLM:
every signal here is a structural/regex fact about real runs, not a
judgment call, so there's nothing for a model to get wrong or invent.
"""

import re

from docx import Document

_BRACKET_STACK_RE = re.compile(r"^\s*[\(\[](.+)[\)\]]\s*$")


def _paragraph_has_hyperlink(paragraph) -> bool:
    return any(
        node.tag.endswith("}hyperlink") for node in paragraph._p.iter()
    )


def _paragraph_run_texts(paragraph) -> list[str]:
    """All literal text runs in a paragraph, INCLUDING ones inside a
    hyperlink (paragraph.runs alone would silently omit those — the
    same quirk documented in file_parsing.py/docx_writer.py)."""
    return [
        node.text for node in paragraph._p.iter()
        if node.tag.endswith("}t") and node.text
    ]


def detect_title_line_style(source_path: str, titles_by_type: list[tuple[int, str]]) -> dict:
    """
    titles_by_type: [(paragraph_index, entry_type), ...] for EVERY
    experience/project title line in the resume — entry_type is
    "experience" or "project". Pass all of them, not just the ones
    about to be edited.

    Returns a profile PER TYPE, since a resume's own convention
    routinely differs between the two (e.g. projects show a tech stack
    and a GitHub link, experience entries show neither) — using one
    merged profile for both would incorrectly force a tech-stack/link
    segment onto a replaced EXPERIENCE title just because projects
    elsewhere in the same resume happen to have one:
      {
        "project": {"has_tech_stack": True, "bracket_style": "()", "has_link": True},
        "experience": {"has_tech_stack": False, "bracket_style": "()", "has_link": False},
      }
    """
    doc = Document(source_path)
    profile_by_type = {}

    for entry_type in ("project", "experience"):
        indices = [i for i, t in titles_by_type if t == entry_type]
        paragraphs = [doc.paragraphs[i] for i in indices if i < len(doc.paragraphs)]

        stack_count = 0
        link_count = 0
        separators = []

        for p in paragraphs:
            run_texts = _paragraph_run_texts(p)
            if any(_BRACKET_STACK_RE.match(t) for t in run_texts):
                stack_count += 1
                for t in run_texts:
                    m = _BRACKET_STACK_RE.match(t)
                    if m:
                        separators.append("[]" if t.strip().startswith("[") else "()")
            if _paragraph_has_hyperlink(p):
                link_count += 1

        total = max(len(paragraphs), 1)
        profile_by_type[entry_type] = {
            "has_tech_stack": stack_count >= max(1, total // 2),
            "bracket_style": (
                max(set(separators), key=separators.count) if separators else "()"
            ),
            "has_link": link_count >= max(1, total // 2),
        }

    return profile_by_type


def detect_bullet_bold_style(source_path: str, bullet_indices: list[int]) -> bool:
    """
    Returns True if this resume's bullets bold specific terms inline
    (mixed bold/non-bold runs within the SAME bullet paragraph) — as
    opposed to a bullet being entirely bold or entirely plain, which
    isn't the "bolded keywords" pattern this is looking for.

    bullet_indices: every bullet paragraph index across the resume, so
    one atypical bullet can't flip the detected convention either way.
    """
    doc = Document(source_path)
    mixed_bold_count = 0
    checked = 0

    for i in bullet_indices:
        if i >= len(doc.paragraphs):
            continue
        runs = [r for r in doc.paragraphs[i].runs if r.text.strip()]
        if len(runs) < 2:
            continue
        checked += 1
        bold_flags = {bool(r.bold) for r in runs}
        if len(bold_flags) > 1:  # both True and False present
            mixed_bold_count += 1

    if checked == 0:
        return False
    return mixed_bold_count >= max(1, checked // 3)